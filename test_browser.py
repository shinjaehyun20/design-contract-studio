"""Regression checks for the public presentation editor."""
import json
import os
from pathlib import Path

from playwright.sync_api import expect, sync_playwright

OUT = Path(os.environ.get("QA_OUTPUT", "/tmp/dcs-quality"))
OUT.mkdir(parents=True, exist_ok=True)
BASE = os.environ.get("QA_URL", "http://127.0.0.1:8765/")


def exercise_editing(browser):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Inches, Pt

    page = browser.new_page(viewport={"width": 1440, "height": 900}, accept_downloads=True)
    checks = []
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    try:
        page.goto(BASE, wait_until="networkidle", timeout=90000)
        page.wait_for_function("window.fileToolStatus?.ready === true", timeout=180000)
        before = page.evaluate("window.CanvasStudio.getProject().slides")
        page.locator("#showcaseLoad").click()
        page.wait_for_function("['contract-loaded','failed'].includes(window.CanvasStudio.getShowcaseEvidence().phase)", timeout=120000)
        evidence = page.evaluate("window.CanvasStudio.getShowcaseEvidence()")
        assert not evidence["error"], evidence
        assert page.evaluate("window.CanvasStudio.getProject().slides") == before, "Loading rules changed the slide"
        checks.append("Loading DESIGN.md leaves slide content unchanged")
        page.locator("#showcaseApply").click()
        page.wait_for_function("window.CanvasStudio.getShowcaseEvidence().phase === 'applied'")
        evidence = page.evaluate("window.CanvasStudio.getShowcaseEvidence()")
        assert evidence["apply"]["changedObjects"] == 3, evidence
        checks.append("Apply changes the three targeted sample objects")

        page.locator('.canvas-object[data-id="headline"]').click()
        headline = "window.CanvasStudio.getProject().slides[0].objects.find(o=>o.id==='headline').content"
        original = page.evaluate(headline)
        # Selection fields synchronize on a deferred pointerup callback.
        # Read the model and wait for the visible editor before changing it.
        expect(page.locator("#publicText")).to_have_value(original)
        edited = "A better workday.\nA clearer story."
        page.locator("#publicText").fill(edited)
        page.locator("#publicUpdate").click()
        assert page.evaluate(headline) == edited
        page.locator("#publicUndo").click()
        assert page.evaluate(headline) == original, "Undo did not restore the selected text"
        page.locator("#publicRedo").click()
        assert page.evaluate(headline) == edited, "Redo did not restore the edit"
        checks.append("Text update, Undo, and Redo preserve the expected content")

        page.locator(".file-note summary").click()
        with page.expect_download(timeout=120000) as download:
            page.locator("#publicSaveJson").click()
        json_path = OUT / "saved-project.json"
        download.value.save_as(json_path)
        saved = json.loads(json_path.read_text(encoding="utf-8"))
        assert saved["slides"] == page.evaluate("window.CanvasStudio.getProject().slides")
        checks.append("Downloaded project JSON matches the current slides")

        with page.expect_download(timeout=120000) as download:
            page.locator("#publicExport").click()
        pptx_path = OUT / "edited-northstar.pptx"
        download.value.save_as(pptx_path)
        native = Presentation(str(pptx_path))
        assert len(native.slides) == 1
        assert edited in [shape.text.replace("\v", "\n") for shape in native.slides[0].shapes if shape.has_text_frame]
        checks.append("Downloaded PPTX contains the edited, editable text")
        page.locator("#showcaseRoundTrip").click()
        page.wait_for_function("['pass','failed'].includes(window.CanvasStudio.getShowcaseEvidence().phase)", timeout=120000)
        one_slide = page.evaluate("window.CanvasStudio.getShowcaseEvidence().roundTrip")
        assert one_slide and one_slide.get("pass"), one_slide
        page.screenshot(path=str(OUT / "editing-roundtrip.png"), full_page=True)
        checks.append("Export and reopen succeeds for the edited one-slide presentation")

        page.locator("#showcaseReset").click()
        transfer = page.evaluate_handle("""data => {
            const transfer=new DataTransfer();
            transfer.items.add(new File([data], 'saved-project.json', {type:'application/json'}));
            return transfer;
        }""", json_path.read_text(encoding="utf-8"))
        page.dispatch_event("#canvasViewport", "drop", {"dataTransfer": transfer})
        page.wait_for_function("window.CanvasStudio.getProject().slides[0].objects.find(o=>o.id==='headline')?.content === 'A better workday.\\nA clearer story.'")
        assert page.evaluate("window.CanvasStudio.getProject().slides") == saved["slides"]
        checks.append("Dropping the saved JSON restores all saved slide objects")

        source = Presentation()
        source.slide_width = Inches(13.333333)
        source.slide_height = Inches(7.5)
        for number in range(1, 4):
            slide = source.slides.add_slide(source.slide_layouts[6])
            title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
            title.text = f"Northstar {number}\n브랜드 규칙 {number}"
            for paragraph in title.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(28)
                    run.font.color.rgb = RGBColor.from_string("1D5D46")
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9), Inches(1), Inches(2), Inches(2))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string("D5DBC1")
            slide.notes_slide.notes_text_frame.text = f"Northstar page {number}"
        source_path = OUT / "three-slide-source.pptx"
        source.save(str(source_path))
        with page.expect_file_chooser() as chooser:
            page.locator("#publicOpen").click()
        chooser.value.set_files(source_path)
        page.wait_for_function("window.CanvasStudio.getProject().slides.length === 3", timeout=120000)
        page.locator("#publicSlide").select_option("2")
        assert page.evaluate("window.CanvasStudio.getEditorState().activeSlide") == 2
        checks.append("Open PPTX imports three slides and slide navigation selects the third")
        with page.expect_download(timeout=120000) as download:
            page.locator("#publicExport").click()
        multi_path = OUT / "three-slide-export.pptx"
        download.value.save_as(multi_path)
        exported = Presentation(str(multi_path))
        assert len(exported.slides) == 3
        for number, slide in enumerate(exported.slides, start=1):
            assert f"Northstar {number}\n브랜드 규칙 {number}" in [shape.text.replace("\v", "\n") for shape in slide.shapes if shape.has_text_frame]
        checks.append("Three-slide PPTX export preserves every page and its Korean/English text")
        page.locator("#showcaseRoundTrip").click()
        page.wait_for_function("['pass','failed'].includes(window.CanvasStudio.getShowcaseEvidence().phase)", timeout=120000)
        multi = page.evaluate("window.CanvasStudio.getShowcaseEvidence().roundTrip")
        assert multi and multi.get("pass") and multi["slides"] == 3, multi
        checks.append("Export and reopen succeeds for the three-slide presentation")
        page.screenshot(path=str(OUT / "three-slide-roundtrip.png"), full_page=True)
        assert not errors, errors
        return {"checks": checks, "oneSlide": one_slide, "threeSlides": multi, "pageErrors": errors}
    finally:
        (OUT / "interaction-progress.json").write_text(json.dumps({"checks": checks, "pageErrors": errors}, ensure_ascii=False, indent=2), encoding="utf-8")
        page.close()


def main():
    findings = []
    results = []
    interaction = None
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=os.environ.get("QA_CHROMIUM"))
        try:
            for width, height in [(1440, 900), (768, 1024), (390, 844), (320, 800)]:
                page = browser.new_page(viewport={"width": width, "height": height})
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                try:
                    page.goto(BASE, wait_until="networkidle", timeout=90000)
                    page.wait_for_function("Boolean(window.CanvasStudio && document.querySelector('#publicSlide'))")
                    page.locator("#shapeMenuButton").wait_for(state="visible")
                    page.screenshot(path=str(OUT / f"workspace-{width}.png"), full_page=True)
                    metrics = page.evaluate("""() => {
                      const rect = el => { const r=el.getBoundingClientRect();return {x:r.x,y:r.y,width:r.width,height:r.height,right:r.right,bottom:r.bottom}; };
                      const bar=document.querySelector('#editModeBar');
                      const shape=document.querySelector('#shapeMenuButton');
                      const range=document.createRange();range.selectNodeContents(shape);
                      const sr=shape.getBoundingClientRect(),tr=range.getBoundingClientRect();
                      const controls=[...bar.querySelectorAll('button,select')].filter(e=>e.getClientRects().length);
                      const overlapPairs=[];
                      for(let i=0;i<controls.length;i++)for(let j=i+1;j<controls.length;j++){
                        const a=controls[i].getBoundingClientRect(),b=controls[j].getBoundingClientRect();
                        if(Math.min(a.right,b.right)>Math.max(a.left,b.left)+1 && Math.min(a.bottom,b.bottom)>Math.max(a.top,b.top)+1)
                          overlapPairs.push([controls[i].id||controls[i].dataset.tool,controls[j].id||controls[j].dataset.tool]);
                      }
                      const headerClipped=[...document.querySelectorAll('.header-actions button')].some(el=>{const r=el.getBoundingClientRect();return r.left<0||r.right>innerWidth+1;});
                      return {title:document.title,viewport:innerWidth,documentWidth:document.documentElement.scrollWidth,
                        bar:rect(bar),shape:rect(shape),overlapPairs,headerClipped,
                        labelContained:tr.left>=sr.left-1 && tr.right<=sr.right+1,
                        controls:controls.map(e=>({id:e.id||e.dataset.tool,...rect(e)}))};
                    }""")
                    failures = []
                    if metrics["documentWidth"] > width + 1:
                        failures.append("page has horizontal overflow")
                    if metrics["headerClipped"]:
                        failures.append("header actions extend beyond the viewport")
                    if metrics["overlapPairs"]:
                        failures.append("editing controls overlap: " + str(metrics["overlapPairs"]))
                    if not metrics["labelContained"]:
                        failures.append("Shapes label extends beyond its button")
                    page.locator("#shapeMenuButton").click()
                    page.locator("#shapePopover").wait_for(state="visible")
                    metrics["shapeMenuOpened"] = True
                    if errors:
                        failures.append("page errors: " + "; ".join(errors))
                    metrics["failures"] = failures
                    results.append(metrics)
                    findings.extend(f"{width}px: {failure}" for failure in failures)
                except Exception as error:
                    findings.append(f"{width}px: {type(error).__name__}: {error}")
                    page.screenshot(path=str(OUT / f"error-{width}.png"), full_page=True)
                finally:
                    page.close()
            try:
                interaction = exercise_editing(browser)
            except Exception as error:
                findings.append(f"editing flow: {type(error).__name__}: {error}")
        finally:
            browser.close()
    report = {"url": BASE, "results": results, "interaction": interaction, "findings": findings}
    (OUT / "browser-results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if findings:
        raise AssertionError("\n".join(findings))


if __name__ == "__main__":
    main()
