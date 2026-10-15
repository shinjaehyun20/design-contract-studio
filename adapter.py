from __future__ import annotations

import base64, hashlib, io, json, mimetypes, posixpath, re, time, uuid, zipfile

from pathlib import Path

from typing import Any, Iterable

from xml.etree import ElementTree as ET

import yaml

from defusedxml import ElementTree as SafeET

from pptx import Presentation

from pptx.dml.color import RGBColor

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, MSO_SHAPE_TYPE

from pptx.enum.text import PP_ALIGN

from pptx.util import Inches, Pt

from pptx.oxml.xmlchemy import OxmlElement

MAX_XML_PART = 10 * 1024 * 1024

P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"

PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

P14_SECTION_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"

PRESETS = {
    "wide": {"label": "Wide 16:9", "width": 960, "height": 540},
    "standard": {"label": "Standard 4:3", "width": 960, "height": 720},
    "a4-landscape": {"label": "A4 landscape", "width": 960, "height": 679},
    "a4-portrait": {"label": "A4 portrait", "width": 679, "height": 960},
}

def uid(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:12]}"

def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]

def read_xml_part(archive: zipfile.ZipFile, name: str) -> Any:
    info = archive.getinfo(name)
    if info.file_size > MAX_XML_PART:
        raise ValueError(f"OOXML part is too large: {name}")
    return SafeET.fromstring(archive.read(info))

def enum_name(value: Any) -> str:
    return str(getattr(value, "name", "") or str(value).split(" ", 1)[0]).lower()

def placeholder_metadata(shape: Any) -> dict[str, Any]:
    try:
        if not shape.is_placeholder:
            return {}
        placeholder = shape.placeholder_format
        return {"placeholder": {"type": enum_name(placeholder.type), "index": int(placeholder.idx)}}
    except Exception:
        return {}

def theme_metadata(root: Any, theme_id: str, source_id: str) -> dict[str, Any]:
    scheme = next((node for node in root.iter() if local_name(node.tag) == "clrScheme"), None)
    colors: dict[str, str] = {}
    if scheme is not None:
        for color in list(scheme):
            value = next(iter(color), None)
            if value is not None:
                colors[local_name(color.tag)] = value.attrib.get("val") or value.attrib.get("lastClr") or ""
    fonts: dict[str, str] = {}
    for family in ("majorFont", "minorFont"):
        node = next((item for item in root.iter() if local_name(item.tag) == family), None)
        if node is None:
            continue
        prefix = "major" if family == "majorFont" else "minor"
        for child in list(node):
            name = local_name(child.tag)
            if name in {"latin", "ea", "cs"}:
                fonts[f"{prefix}{name.title()}"] = child.attrib.get("typeface", "")
    return {
        "id": theme_id,
        "sourceId": source_id,
        "name": root.attrib.get("name") or Path(source_id).stem,
        "colors": colors,
        "fonts": fonts,
    }

def inspect_pptx_package(data: bytes) -> dict[str, Any]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = set(archive.namelist())
        presentation = read_xml_part(archive, "ppt/presentation.xml")
        slide_list = next((node for node in list(presentation) if local_name(node.tag) == "sldIdLst"), None)
        slide_ids = [child.attrib.get("id", "") for child in list(slide_list or []) if local_name(child.tag) == "sldId"]
        slide_indexes = {slide_id: index for index, slide_id in enumerate(slide_ids)}

        sections: list[dict[str, Any]] = []
        for section in (node for node in presentation.iter() if local_name(node.tag) == "section"):
            member_ids = [node.attrib.get("id", "") for node in section.iter() if local_name(node.tag) == "sldId"]
            indexes = [slide_indexes[slide_id] for slide_id in member_ids if slide_id in slide_indexes]
            sections.append({
                "sourceId": section.attrib.get("id", ""),
                "name": section.attrib.get("name") or f"Section {len(sections) + 1}",
                "slideIndexes": indexes,
            })

        themes: list[dict[str, Any]] = []
        theme_ids_by_path: dict[str, str] = {}
        for source_id in sorted(name for name in names if name.startswith("ppt/theme/theme") and name.endswith(".xml")):
            theme_id = uid("theme")
            theme_ids_by_path[source_id] = theme_id
            themes.append(theme_metadata(read_xml_part(archive, source_id), theme_id, source_id))

        master_theme_ids: dict[str, str] = {}
        master_parts = sorted(name for name in names if name.startswith("ppt/slideMasters/slideMaster") and name.endswith(".xml"))
        for master_part in master_parts:
            rels_name = posixpath.join(posixpath.dirname(master_part), "_rels", f"{posixpath.basename(master_part)}.rels")
            if rels_name not in names:
                continue
            rels = read_xml_part(archive, rels_name)
            relationship = next((node for node in rels if node.attrib.get("Type", "").endswith("/theme")), None)
            if relationship is None:
                continue
            target = posixpath.normpath(posixpath.join(posixpath.dirname(master_part), relationship.attrib.get("Target", "")))
            if target in theme_ids_by_path:
                master_theme_ids[master_part] = theme_ids_by_path[target]

    return {"sections": sections, "themes": themes, "masterThemeIds": master_theme_ids}

def parse_design_md(data: bytes, filename: str) -> dict[str, Any]:
    if len(data) > 2 * 1024 * 1024:
        raise ValueError("DESIGN.md must be 2 MB or smaller")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("DESIGN.md must be UTF-8 text") from exc
    if "\x00" in text:
        raise ValueError("DESIGN.md must not contain NUL bytes")
    if not text.strip():
        raise ValueError("DESIGN.md is empty")

    frontmatter: dict[str, Any] = {}
    body_text = text
    lines = text.splitlines()
    if lines and lines[0].strip() == "---":
        closing = next((index for index, line in enumerate(lines[1:], 1) if line.strip() == "---"), None)
        if closing is None:
            raise ValueError("DESIGN.md YAML frontmatter is not closed")
        yaml_source = "\n".join(lines[1:closing])
        try:
            if any(isinstance(token, (yaml.tokens.AnchorToken, yaml.tokens.AliasToken)) for token in yaml.scan(yaml_source)):
                raise ValueError("DESIGN.md YAML anchors and aliases are not allowed")
            loaded = yaml.safe_load(yaml_source) or {}
        except ValueError:
            raise
        except yaml.YAMLError as exc:
            raise ValueError(f"DESIGN.md YAML frontmatter is invalid: {exc}") from exc
        if not isinstance(loaded, dict):
            raise ValueError("DESIGN.md YAML frontmatter must be a mapping")
        frontmatter = loaded
        body_text = "\n".join(lines[closing + 1:])

    def json_safe(value: Any) -> Any:
        if isinstance(value, dict):
            return {str(key): json_safe(item) for key, item in value.items()}
        if isinstance(value, (list, tuple)):
            return [json_safe(item) for item in value]
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        return str(value)

    frontmatter = json_safe(frontmatter)

    headings: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    rules: list[dict[str, Any]] = []
    automation: list[str] = []
    colors: list[dict[str, Any]] = []
    typography: list[dict[str, Any]] = []
    spacing: list[dict[str, Any]] = []
    seen_colors: set[tuple[str, str]] = set()
    seen_type: set[str] = set()
    seen_spacing: set[str] = set()
    section_title = ""
    component_section = False
    current_component: dict[str, Any] | None = None
    warnings: list[str] = []

    def add_color(name: str, value: Any) -> None:
        normalized = str(value).upper()
        if not re.fullmatch(r"#[0-9A-F]{6,8}", normalized):
            warnings.append(f"Invalid color token ignored: {name}")
            return
        key = (name.lower(), normalized)
        if key not in seen_colors:
            seen_colors.add(key)
            colors.append({"name": name, "value": normalized})

    front_colors = frontmatter.get("colors") or {}
    if isinstance(front_colors, dict):
        for name, value in front_colors.items():
            add_color(str(name), value)

    front_typography = frontmatter.get("typography") or {}
    if isinstance(front_typography, dict):
        for name, value in front_typography.items():
            properties = value if isinstance(value, dict) else {"fontFamily": value}
            family = properties.get("fontFamily")
            if not family:
                warnings.append(f"Typography token has no fontFamily: {name}")
                continue
            family_text = str(family)
            if family_text not in seen_type:
                seen_type.add(family_text)
                typography.append({"name": str(name), "value": family_text, "properties": properties})

    front_spacing = frontmatter.get("spacing") or {}
    if isinstance(front_spacing, dict):
        for name, raw_value in front_spacing.items():
            match = re.fullmatch(r"(\d+(?:\.\d+)?)px", str(raw_value).strip(), re.IGNORECASE)
            if not match:
                warnings.append(f"Invalid spacing token ignored: {name}")
                continue
            key = str(name).lower()
            if key not in seen_spacing:
                seen_spacing.add(key)
                number: int | float = float(match.group(1)) if "." in match.group(1) else int(match.group(1))
                spacing.append({"name": str(name), "value": number, "unit": "px"})

    front_components = frontmatter.get("components") or {}
    if isinstance(front_components, dict):
        for name, properties in front_components.items():
            props = properties if isinstance(properties, dict) else {"value": properties}
            components.append({"id": f"component-{len(components)+1}", "name": str(name), "description": str(props.get("description", "")), "properties": props})

    for raw_line in body_text.splitlines():
        line = raw_line.strip()
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            level, title = len(heading.group(1)), heading.group(2).strip()
            headings.append({"level": level, "title": title})
            if level <= 2:
                section_title = title
                component_section = any(word in title.lower() for word in ("component", "컴포넌트"))
            current_component = None
            if level >= 3 and component_section:
                current_component = {"id": f"component-{len(components)+1}", "name": title, "description": ""}
                components.append(current_component)
            continue

        css_color = re.search(r"(--[\w-]+)\s*:\s*(#[0-9a-fA-F]{6,8})\b", line)
        table_color = re.match(r"^\|\s*([^|]+?)\s*\|\s*(#[0-9a-fA-F]{6,8})\s*\|?", line)
        generic_color = re.search(r"([A-Za-z가-힣][\w가-힣 .-]{1,40})\s*[:=]\s*(#[0-9a-fA-F]{6,8})\b", line)
        color_match = css_color or table_color or generic_color
        if color_match:
            name = color_match.group(1).strip(" -|`").replace("--", "")
            value = color_match.group(2).upper()
            key = (name.lower(), value)
            if key not in seen_colors:
                seen_colors.add(key)
                colors.append({"name": name, "value": value})

        family = re.search(r"font-family\s*:\s*([^;|]+)", line, re.IGNORECASE)
        if family:
            value = family.group(1).strip()
            if value not in seen_type:
                seen_type.add(value)
                typography.append({"name": "font-family", "value": value})

        space = re.search(r"(--(?:space|spacing|gap|padding)-[\w-]+)\s*:\s*(\d+(?:\.\d+)?)px", line, re.IGNORECASE)
        if space and space.group(1).lower() not in seen_spacing:
            seen_spacing.add(space.group(1).lower())
            value: int | float = float(space.group(2)) if "." in space.group(2) else int(space.group(2))
            spacing.append({"name": space.group(1).removeprefix("--"), "value": value, "unit": "px"})

        bullet = re.match(r"^[-*+]\s+(.+)$", line)
        if bullet:
            statement = bullet.group(1).strip()
            context = section_title.lower()
            if any(word in context for word in ("rule", "규칙", "금지", "검증", "must", "avoid")):
                lowered = statement.lower()
                level = "avoid" if any(word in lowered for word in ("금지", "avoid", "do not", "don't")) else "must" if any(word in lowered for word in ("must", "필수", "해야", "한다")) else "guidance"
                rules.append({"id": f"rule-{len(rules)+1}", "level": level, "text": statement, "category": section_title})
            if any(word in context for word in ("automation", "자동화", "runtime", "런타임")):
                automation.append(statement)
            if current_component and not current_component["description"]:
                current_component["description"] = statement
        elif line and current_component and not current_component["description"]:
            current_component["description"] = line

    title = str(frontmatter.get("name") or next((item["title"] for item in headings if item["level"] == 1), Path(filename).stem))
    extensions = {key: value for key, value in frontmatter.items() if key.startswith("x-")}
    return {
        "schema": "design-md/1", "name": title,
        "source": {
            "format": "design-md", "fileName": Path(filename).name,
            "sha256": hashlib.sha256(data).hexdigest(), "byteLength": len(data),
            "parserVersion": "design-md-parser/1.1.0",
        },
        "tokens": {"colors": colors, "typography": typography, "spacing": spacing},
        "rules": rules, "components": components,
        "automation": {"hints": automation}, "outline": headings,
        "extensions": extensions,
        "spec": {"version": frontmatter.get("version"), "description": frontmatter.get("description"), "frontmatter": frontmatter},
        "report": {"errors": [], "warnings": warnings, "summary": {"errors": 0, "warnings": len(warnings)}},
        "sourceText": text,
    }

def rgb_hex(color: Any, fallback: str) -> str:
    try:
        rgb = color.rgb
        if rgb is not None:
            return f"#{str(rgb)}"
    except Exception:
        pass
    return fallback

def fill_hex(shape: Any, fallback: str = "transparent") -> str:
    try:
        fill = shape.fill
        if fill.type is None:
            return fallback
        return rgb_hex(fill.fore_color, fallback)
    except Exception:
        return fallback

def line_hex(shape: Any, fallback: str = "transparent") -> str:
    try:
        if shape.line.fill.type is None:
            return fallback
        return rgb_hex(shape.line.color, fallback)
    except Exception:
        return fallback

def line_width(shape: Any, scale: float = 4/3) -> float:
    try:
        return max(0.0, round(float(shape.line.width.pt) * scale, 2))
    except Exception:
        return 0.0

def first_run(shape: Any) -> Any | None:
    try:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text:
                    return run
    except Exception:
        return None
    return None

def text_style(shape: Any, scale: float) -> dict[str, Any]:
    run = first_run(shape)
    font = getattr(run, "font", None)
    size = 18
    try:
        if font and font.size:
            size = max(6, round(font.size.pt * scale, 1))
    except Exception:
        pass
    name = "Pretendard, Noto Sans KR, sans-serif"
    try:
        if font and font.name:
            name = f"{font.name}, Pretendard, Noto Sans KR, sans-serif"
    except Exception:
        pass
    color = "#222320"
    try:
        if font:
            color = rgb_hex(font.color, color)
    except Exception:
        pass
    align = "left"
    try:
        value = shape.text_frame.paragraphs[0].alignment
        align = {1: "left", 2: "center", 3: "right"}.get(int(value), "left")
    except Exception:
        pass
    return {
        "fontFamily": name,
        "fontSize": size,
        "fontWeight": 700 if bool(getattr(font, "bold", False)) else 400,
        "fontStyle": "italic" if bool(getattr(font, "italic", False)) else "normal",
        "color": color,
        "textAlign": align,
        "lineHeight": float(shape.text_frame.paragraphs[0].line_spacing) if isinstance(shape.text_frame.paragraphs[0].line_spacing, float) else 1.0,
        "background": fill_hex(shape),
        "borderColor": line_hex(shape),
        "borderWidth": line_width(shape, scale),
        "radius": 0,
        "padding": round(float(shape.text_frame.margin_left) / 12700 * scale, 2),
        "opacity": 1,
    }

def shape_name(shape: Any, fallback: str) -> str:
    value = str(getattr(shape, "name", "") or "").strip()
    return value or fallback

def shape_source_metadata(shape: Any) -> dict[str, Any]:
    return {
        "source": {
            "officeName": shape_name(shape, "Office shape"),
            "officeShapeType": enum_name(getattr(shape, "shape_type", "unknown")),
        },
        **placeholder_metadata(shape),
    }

def normalized_text(value: Any) -> str:
    """Normalize PowerPoint soft line breaks for browser-editable text."""
    return str(value or "").replace("\v", "\n").replace("\r\n", "\n").replace("\r", "\n").strip()

def geometry(shape: Any, sx: float, sy: float) -> dict[str, int | float]:
    return {
        "x": round(float(shape.left) * sx),
        "y": round(float(shape.top) * sy),
        "w": max(1, round(float(shape.width) * sx)),
        "h": max(1, round(float(shape.height) * sy)),
        "rotation": round(float(getattr(shape, "rotation", 0) or 0), 2),
    }

def extract_notes(slide: Any) -> str:
    try:
        texts: list[str] = []
        for shape in slide.notes_slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = normalized_text(shape.text)
                if text and text not in texts:
                    texts.append(text)
        return "\n".join(texts)
    except Exception:
        return ""

def table_objects(shape: Any, geom: dict[str, Any], z_start: int, text_scale: float) -> list[dict[str, Any]]:
    table = shape.table
    col_widths = [float(column.width) for column in table.columns]
    row_heights = [float(row.height) for row in table.rows]
    total_w, total_h = max(1.0, sum(col_widths)), max(1.0, sum(row_heights))
    objects: list[dict[str, Any]] = []
    y = geom["y"]
    for row_index, row in enumerate(table.rows):
        height = geom["h"] * row_heights[row_index] / total_h
        x = geom["x"]
        for col_index, cell in enumerate(row.cells):
            width = geom["w"] * col_widths[col_index] / total_w
            text = normalized_text(cell.text)
            cell_shape = {
                "id": uid("text"), "type": "text",
                "name": f"{shape_name(shape, 'Table')} R{row_index + 1}C{col_index + 1}",
                "content": text, "x": round(x), "y": round(y), "w": max(2, round(width)), "h": max(2, round(height)),
                "z": z_start + len(objects), "rotation": 0,
                "style": {
                    "fontFamily": "Pretendard, Noto Sans KR, sans-serif", "fontSize": max(7, round(12 * text_scale, 1)),
                    "fontWeight": 700 if row_index == 0 else 400, "color": "#222320", "textAlign": "left",
                    "lineHeight": float(shape.text_frame.paragraphs[0].line_spacing) if isinstance(shape.text_frame.paragraphs[0].line_spacing, float) else 1.25, "background": "#f3f2ee" if row_index == 0 else "#ffffff",
                    "borderColor": "#c8c7c1", "borderWidth": 1, "radius": 0, "padding": 4, "opacity": 1,
                },
                "role": "table-cell",
            }
            objects.append(cell_shape)
            x += width
        y += height
    return objects

def iter_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)
        else:
            yield shape

def native_shape_kind(shape: Any) -> str:
    mapping = {
        MSO_AUTO_SHAPE_TYPE.RECTANGLE: 'rectangle',
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE: 'rounded',
        MSO_AUTO_SHAPE_TYPE.OVAL: 'oval',
        MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE: 'triangle',
        MSO_AUTO_SHAPE_TYPE.DIAMOND: 'diamond',
        MSO_AUTO_SHAPE_TYPE.HEXAGON: 'hexagon',
        MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW: 'right-arrow',
        MSO_AUTO_SHAPE_TYPE.CHEVRON: 'chevron',
        MSO_AUTO_SHAPE_TYPE.RECTANGULAR_CALLOUT: 'callout',
    }
    try:
        return mapping.get(shape.auto_shape_type, 'rectangle')
    except (AttributeError, ValueError):
        return 'rectangle'

def convert_shapes(shapes: Iterable[Any], sx: float, sy: float, text_scale: float) -> list[dict[str, Any]]:
    objects: list[dict[str, Any]] = []
    for z, shape in enumerate(iter_shapes(shapes), start=1):
        geom = geometry(shape, sx, sy)
        try:
            if getattr(shape, "has_table", False):
                objects.extend(table_objects(shape, geom, z * 100, text_scale))
                continue
        except Exception:
            pass
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
            image = shape.image
            content_type = image.content_type or mimetypes.guess_type(image.filename)[0] or "image/png"
            objects.append({
                "id": uid("image"), "type": "image", "name": shape_name(shape, image.filename or "Picture"),
                **geom, "z": z, "src": f"data:{content_type};base64,{base64.b64encode(image.blob).decode('ascii')}",
                "alt": shape_name(shape, "Imported picture"), "style": {"objectFit": "contain", "opacity": 1},
                **shape_source_metadata(shape),
            })
            continue
        text = ""
        try:
            if getattr(shape, "has_text_frame", False):
                text = normalized_text(shape.text)
        except Exception:
            pass
        if text:
            objects.append({
                "id": uid("text"), "type": "text", "name": shape_name(shape, f"Text {z}"), "content": text,
                **geom, "z": z, "style": text_style(shape, text_scale),
                **shape_source_metadata(shape),
            })
            continue
        unsupported = getattr(shape, "has_chart", False)
        objects.append({
            "id": uid("shape"), "type": "shape", "name": shape_name(shape, f"Shape {z}"), **geom, "z": z,
            "kind": native_shape_kind(shape),
            "style": {"fill": fill_hex(shape), "borderColor": line_hex(shape), "borderWidth": line_width(shape, text_scale), "radius": 999 if native_shape_kind(shape)=="oval" else 0, "opacity": 1},
            **shape_source_metadata(shape),
            **({"role": "unsupported-chart"} if unsupported else {}),
        })
        if unsupported:
            objects.append({
                "id": uid("text"), "type": "text", "name": f"{shape_name(shape, 'Chart')} label", "content": "Chart · 원본 PPTX 참조",
                **geom, "z": z + 0.1,
                "style": {"fontFamily": "Pretendard, sans-serif", "fontSize": 12, "fontWeight": 600, "color": "#71726e", "textAlign": "center", "lineHeight": 1.3, "background": "transparent", "borderColor": "transparent", "borderWidth": 0, "radius": 0, "padding": 6, "opacity": 1},
            })
    return objects

def convert_pptx(data: bytes, filename: str) -> dict[str, Any]:
    presentation = Presentation(io.BytesIO(data))
    package = inspect_pptx_package(data)
    ratio = float(presentation.slide_width) / max(1.0, float(presentation.slide_height))
    preset_key = min(PRESETS, key=lambda key: abs(PRESETS[key]["width"] / PRESETS[key]["height"] - ratio))
    preset = PRESETS[preset_key]
    sx = preset["width"] / float(presentation.slide_width)
    sy = preset["height"] / float(presentation.slide_height)
    text_scale = min(sx, sy) * 12700  # EMU scale to a 96-ish DPI canvas, normalized below
    text_scale = max(0.65, min(1.35, text_scale))
    slides: list[dict[str, Any]] = []
    masters: list[dict[str, Any]] = []
    master_ids: dict[str, str] = {}
    layout_ids: dict[str, str] = {}

    section_specs = package["sections"] or [{"sourceId": "", "name": "Imported deck", "slideIndexes": list(range(len(presentation.slides)))}]
    sections: list[dict[str, Any]] = []
    section_by_slide: dict[int, str] = {}
    for section_index, spec in enumerate(section_specs):
        section_id = uid("section")
        sections.append({
            "id": section_id,
            "sourceId": spec.get("sourceId") or "",
            "name": spec.get("name") or f"Section {section_index + 1}",
            "collapsed": False,
        })
        for slide_index in spec.get("slideIndexes", []):
            section_by_slide.setdefault(int(slide_index), section_id)
    unassigned = [index for index in range(len(presentation.slides)) if index not in section_by_slide]
    if unassigned:
        section_id = uid("section")
        sections.append({"id": section_id, "sourceId": "", "name": "Unsectioned", "collapsed": False})
        section_by_slide.update({index: section_id for index in unassigned})

    for master_index, master in enumerate(presentation.slide_masters):
        master_id = uid("master")
        master_part = str(master.part.partname).lstrip("/")
        master_ids[master_part] = master_id
        layouts: list[dict[str, Any]] = []
        for layout_index, layout in enumerate(master.slide_layouts):
            layout_id = uid("layout")
            layout_part = str(layout.part.partname).lstrip("/")
            layout_ids[layout_part] = layout_id
            layouts.append({
                "id": layout_id,
                "sourceId": layout_part,
                "name": layout.name or f"Layout {layout_index + 1}",
                "background": "#ffffff",
                "sizePreset": preset_key, "width": preset["width"], "height": preset["height"],
                "objects": convert_shapes(layout.shapes, sx, sy, text_scale),
            })
        masters.append({
            "id": master_id, "sourceId": master_part, "title": f"Slide Master {master_index + 1}", "background": "#ffffff",
            "sizePreset": preset_key, "width": preset["width"], "height": preset["height"],
            "objects": convert_shapes(master.shapes, sx, sy, text_scale),
            "themeId": package["masterThemeIds"].get(master_part, package["themes"][0]["id"] if package["themes"] else ""),
            "layouts": layouts,
        })

    for slide_index, slide in enumerate(presentation.slides):
        objects = convert_shapes(slide.shapes, sx, sy, text_scale)
        title = slide._element.cSld.get("name") or f"Slide {slide_index + 1}"
        try:
            if slide.shapes.title and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip().splitlines()[0]
        except Exception:
            pass
        master_part = str(slide.slide_layout.slide_master.part.partname).lstrip("/")
        layout_part = str(slide.slide_layout.part.partname).lstrip("/")
        slides.append({
            "id": uid("slide"), "title": title, "background": fill_hex(slide.background, "#ffffff"), "notes": extract_notes(slide),
            "sizePreset": preset_key, "width": preset["width"], "height": preset["height"], "objects": objects,
            "masterId": master_ids.get(master_part, masters[0]["id"] if masters else ""),
            "layoutId": layout_ids.get(layout_part, ""), "layoutName": slide.slide_layout.name or "Layout",
            "sectionId": section_by_slide.get(slide_index, sections[0]["id"]),
        })

    stem = Path(filename).stem or "Imported presentation"
    all_objects = [item for slide in slides for item in slide["objects"]]
    unsupported_charts = [item for item in all_objects if item.get("role") == "unsupported-chart"]
    import_warnings: list[dict[str, str]] = []
    if unsupported_charts:
        import_warnings.append({"code": "chart-placeholder", "severity": "review", "message": f"차트 {len(unsupported_charts)}개는 편집 가능한 차트가 아니라 원본 PPTX 참조 placeholder로 가져왔습니다."})
    if masters:
        import_warnings.append({"code": "master-export-flattening", "severity": "review", "message": "Master/Layout 편집 시 PPTX 내보내기는 슬라이드 객체로 평탄화될 수 있습니다."})
    import_report = {
        "schema": "canvas-import-report/1",
        "sourceSha256": hashlib.sha256(data).hexdigest(),
        "sourceFileName": filename,
        "editableObjects": len(all_objects),
        "unsupportedCharts": len(unsupported_charts),
        "exportReviewRequired": bool(import_warnings),
        "warnings": import_warnings,
    }
    now = time.strftime("%Y-%m-%dT%H:%M:%S%z")
    return {
        "version": 3, "editingMode": "pptx", "id": uid("project"), "name": presentation.core_properties.title or stem, "fileName": filename,
        "createdAt": now, "updatedAt": now, "source": {"type": "pptx", "fileName": filename, "slideWidthEmu": int(presentation.slide_width), "slideHeightEmu": int(presentation.slide_height)}, "importReport": import_report,
        "slides": slides or [{"id": uid("slide"), "title": "Empty slide", "background": "#ffffff", "notes": "", "sizePreset": preset_key, "width": preset["width"], "height": preset["height"], "objects": []}],
        "masters": masters,
        "themes": package["themes"],
        "sections": sections,
        "comments": [], "components": [], "conversation": [{"id": uid("msg"), "role": "assistant", "text": f"{filename}에서 {len(slides)}개 슬라이드와 {sum(len(s['objects']) for s in slides)}개 편집 객체를 가져왔습니다."}],
        "tweaks": {"descriptionWidth": 24, "radius": 4, "density": 100, "accent": "#315ce7"},
    }

def export_color(value: Any, fallback: str = "000000") -> RGBColor:
    match = re.fullmatch(r"#?([0-9a-fA-F]{6})", str(value or ""))
    return RGBColor.from_string(match.group(1).upper() if match else fallback)

def px(value: Any, scale: float = 1.0) -> Any:
    return Inches(max(0.0, float(value or 0) * scale) / 96.0)

def apply_export_fill(fill: Any, value: Any) -> None:
    if not value or str(value).lower() == "transparent":
        fill.background()
        return
    fill.solid()
    fill.fore_color.rgb = export_color(value, "FFFFFF")

def apply_export_line(line: Any, color: Any, width: Any) -> None:
    width_value = max(0.0, float(width or 0))
    if width_value <= 0 or not color or str(color).lower() == "transparent":
        line.fill.background()
        return
    line.color.rgb = export_color(color, "000000")
    line.width = Pt(width_value * 0.75)

def add_export_text(slide: Any, item: dict[str, Any], sx: float, sy: float) -> None:
    style = item.get("style") or {}
    shape = slide.shapes.add_textbox(px(item.get("x"), sx), px(item.get("y"), sy), px(item.get("w"), sx), px(item.get("h"), sy))
    shape.name = str(item.get("name") or shape.name)
    shape.rotation = float(item.get("rotation") or 0)
    apply_export_fill(shape.fill, style.get("background"))
    apply_export_line(shape.line, style.get("borderColor"), style.get("borderWidth"))
    frame = shape.text_frame
    frame.clear()
    padding = max(0.0, float(style.get("padding") or 0))
    frame.margin_left = frame.margin_right = px(padding, sx)
    frame.margin_top = frame.margin_bottom = px(padding, sy)
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = str(item.get("content") or "")
    paragraph.alignment = {
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }.get(str(style.get("textAlign") or "left"), PP_ALIGN.LEFT)
    try:
        paragraph.line_spacing = float(style.get("lineHeight") or 1.25)
    except Exception:
        pass
    for run in paragraph.runs:
        font = run.font
        font.name = str(style.get("fontFamily") or "Pretendard").split(",")[0].strip()
        font.size = Pt(max(4.5, float(style.get("fontSize") or 14) * min(sx, sy) * 0.75))
        font.bold = float(style.get("fontWeight") or 400) >= 600
        font.italic = str(style.get("fontStyle") or "normal") == "italic"
        font.color.rgb = export_color(style.get("color"), "222320")
        if style.get("strike"):
            run._r.get_or_add_rPr().set("strike", "sngStrike")

def add_export_shape(slide: Any, item: dict[str, Any], sx: float, sy: float) -> None:
    style = item.get("style") or {}
    shape_types = {
        "rectangle": MSO_AUTO_SHAPE_TYPE.RECTANGLE, "rounded": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        "oval": MSO_AUTO_SHAPE_TYPE.OVAL, "triangle": MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
        "diamond": MSO_AUTO_SHAPE_TYPE.DIAMOND, "hexagon": MSO_AUTO_SHAPE_TYPE.HEXAGON,
        "right-arrow": MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, "chevron": MSO_AUTO_SHAPE_TYPE.CHEVRON,
        "callout": MSO_AUTO_SHAPE_TYPE.RECTANGULAR_CALLOUT,
    }
    kind = shape_types.get(item.get("kind"))
    if kind is None:
        kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if float(style.get("radius") or 0) >= 8 else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, px(item.get("x"), sx), px(item.get("y"), sy), px(item.get("w"), sx), px(item.get("h"), sy))
    # Canvas shapes have no theme effects. An explicit empty list prevents
    # PowerPoint's default shape style from adding a shadow on native reopen.
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    shape.name = str(item.get("name") or shape.name)
    shape.rotation = float(item.get("rotation") or 0)
    apply_export_fill(shape.fill, style.get("fill"))
    apply_export_line(shape.line, style.get("borderColor"), style.get("borderWidth"))

def add_export_image(slide: Any, item: dict[str, Any], sx: float, sy: float) -> None:
    source = str(item.get("src") or "")
    if not source.startswith("data:") or "," not in source:
        return
    try:
        payload = base64.b64decode(source.split(",", 1)[1])
        picture = slide.shapes.add_picture(io.BytesIO(payload), px(item.get("x"), sx), px(item.get("y"), sy), px(item.get("w"), sx), px(item.get("h"), sy))
        picture.name = str(item.get("name") or picture.name)
        style = item.get("style") or {}
        zoom = max(1.0, float(style.get("cropZoom") or 1))
        if zoom > 1:
            crop = 1 - 1 / zoom
            x = min(100.0, max(0.0, float(style.get("cropX") or 50))) / 100
            y = min(100.0, max(0.0, float(style.get("cropY") or 50))) / 100
            picture.crop_left, picture.crop_right = crop * x, crop * (1 - x)
            picture.crop_top, picture.crop_bottom = crop * y, crop * (1 - y)
    except Exception:
        return

def add_export_path(slide: Any, item: dict[str, Any], sx: float, sy: float) -> None:
    points = item.get("points") or []
    if len(points) < 2:
        return
    first, last = points[0], points[-1]
    x = float(item.get("x") or 0)
    y = float(item.get("y") or 0)
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        px(x + float(first.get("x") or 0), sx),
        px(y + float(first.get("y") or 0), sy),
        px(x + float(last.get("x") or 0), sx),
        px(y + float(last.get("y") or 0), sy),
    )
    style = item.get("style") or {}
    apply_export_line(connector.line, style.get("stroke"), style.get("strokeWidth"))

def inject_pptx_sections(data: bytes, project: dict[str, Any]) -> bytes:
    sections = project.get("sections") or []
    slides = project.get("slides") or []
    if not sections or not slides:
        return data
    source = zipfile.ZipFile(io.BytesIO(data))
    presentation = read_xml_part(source, "ppt/presentation.xml")
    slide_list = next((node for node in list(presentation) if local_name(node.tag) == "sldIdLst"), None)
    slide_ids = [node.attrib.get("id", "") for node in list(slide_list or []) if local_name(node.tag) == "sldId"]
    if len(slide_ids) != len(slides):
        return data

    ext_list = next((node for node in list(presentation) if local_name(node.tag) == "extLst"), None)
    if ext_list is None:
        ext_list = ET.SubElement(presentation, f"{{{PML_NS}}}extLst")
    for ext in list(ext_list):
        if any(local_name(node.tag) == "sectionLst" for node in ext.iter()):
            ext_list.remove(ext)
    ext = ET.SubElement(ext_list, f"{{{PML_NS}}}ext", {"uri": P14_SECTION_URI})
    section_list = ET.SubElement(ext, f"{{{P14_NS}}}sectionLst")
    for index, section in enumerate(sections):
        member_ids = [slide_ids[i] for i, slide in enumerate(slides) if slide.get("sectionId") == section.get("id")]
        if not member_ids:
            continue
        source_id = str(section.get("sourceId") or "")
        section_guid = source_id if re.fullmatch(r"\{[0-9a-fA-F-]{36}\}", source_id) else f"{{{uuid.uuid4()}}}"
        section_node = ET.SubElement(section_list, f"{{{P14_NS}}}section", {
            "name": str(section.get("name") or f"Section {index + 1}"),
            "id": section_guid,
        })
        member_list = ET.SubElement(section_node, f"{{{P14_NS}}}sldIdLst")
        for slide_id in member_ids:
            ET.SubElement(member_list, f"{{{P14_NS}}}sldId", {"id": slide_id})

    ET.register_namespace("p", PML_NS)
    ET.register_namespace("p14", P14_NS)
    presentation_xml = ET.tostring(presentation, encoding="utf-8", xml_declaration=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, presentation_xml if item.filename == "ppt/presentation.xml" else source.read(item.filename))
    return output.getvalue()

def project_to_pptx(project: dict[str, Any]) -> bytes:
    slides = project.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("Project must contain at least one slide")
    first = slides[0]
    page_width = max(1.0, float(first.get("width") or PRESETS["wide"]["width"]))
    page_height = max(1.0, float(first.get("height") or PRESETS["wide"]["height"]))
    presentation = Presentation()
    presentation.core_properties.title = str(project.get("name") or "Presentation")
    presentation.slide_width = px(page_width)
    presentation.slide_height = px(page_height)
    blank = presentation.slide_layouts[6]
    masters_by_id = {master.get("id"): master for master in project.get("masters", [])}
    layouts_by_id = {layout.get("id"): layout for master in project.get("masters", []) for layout in master.get("layouts", []) if isinstance(layout, dict)}
    for slide_data in slides:
        canvas_width = max(1.0, float(slide_data.get("width") or page_width))
        canvas_height = max(1.0, float(slide_data.get("height") or page_height))
        sx, sy = page_width / canvas_width, page_height / canvas_height
        slide = presentation.slides.add_slide(blank)
        slide._element.cSld.set("name", str(slide_data.get("title") or "Slide"))
        if slide_data.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(slide_data["notes"])
        apply_export_fill(slide.background.fill, slide_data.get("background") or "#ffffff")
        master_objects = [item for item in masters_by_id.get(slide_data.get("masterId"), {}).get("objects", []) if not item.get("placeholder")]
        layout_objects = [item for item in layouts_by_id.get(slide_data.get("layoutId"), {}).get("objects", []) if not item.get("placeholder")]
        objects = [*master_objects, *layout_objects, *(slide_data.get("objects") or [])]
        for item in sorted(objects, key=lambda value: float(value.get("z") or 0)):
            kind = item.get("type")
            if kind == "text":
                add_export_text(slide, item, sx, sy)
            elif kind == "shape":
                add_export_shape(slide, item, sx, sy)
            elif kind == "image":
                add_export_image(slide, item, sx, sy)
            elif kind == "path":
                add_export_path(slide, item, sx, sy)
    stream = io.BytesIO()
    presentation.save(stream)
    return inject_pptx_sections(stream.getvalue(), project)

def dispatch_json(raw: str) -> str:
    request = json.loads(raw)
    operation = request.get('operation')
    if operation == 'design':
        design = parse_design_md(base64.b64decode(request['data']), request['filename'])
        result = {'ok': True, 'designSystem': design, 'summary': {
            'colors': len(design['tokens']['colors']), 'rules': len(design['rules']),
            'components': len(design['components'])}}
    elif operation == 'import':
        data = base64.b64decode(request['data'])
        if len(data) > 25 * 1024 * 1024:
            raise ValueError('Choose a PPTX smaller than 25 MB.')
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 5000 or sum(f.file_size for f in archive.infolist()) > 150 * 1024 * 1024:
                raise ValueError('This presentation is too large to open in the browser.')
        project = convert_pptx(data, request['filename'])
        result = {'ok': True, 'project': project, 'summary': {
            'slides': len(project['slides']), 'sections': len(project['sections']),
            'masters': len(project['masters']), 'themes': len(project['themes']),
            'objects': sum(len(s['objects']) for s in project['slides']),
            'importReport': project['importReport']}}
    elif operation == 'export':
        project=request['project']
        if not isinstance(project, dict) or not 1 <= len(project.get('slides', [])) <= 100:
            raise ValueError('A presentation must contain 1 to 100 slides.')
        result = {'ok': True, 'data': base64.b64encode(project_to_pptx(project)).decode('ascii')}
    else:
        raise ValueError('Unknown file operation.')
    return json.dumps(result, ensure_ascii=False)
