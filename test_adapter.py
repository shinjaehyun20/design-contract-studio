import base64, io, json, unittest, zipfile
from pptx import Presentation
from pptx.util import Inches
from adapter import convert_pptx, project_to_pptx, parse_design_md, dispatch_json

class PresentationRoundTrip(unittest.TestCase):
    def test_export_does_not_inherit_theme_shape_effects(self):
        project={'slides':[{'width':960,'height':540,'objects':[{'type':'shape','kind':'rectangle','x':10,'y':20,'w':100,'h':80,'style':{'fill':'#1D5D46'}}]}]}
        native=Presentation(io.BytesIO(project_to_pptx(project)))
        shape=native.slides[0].shapes[0]
        # Without this direct override, the default effectRef resolves to a
        # theme shadow in PowerPoint even though model/roundtrip checks pass.
        self.assertEqual(len(shape._element.xpath('./p:spPr/a:effectLst')),1)
        self.assertEqual(len(shape._element.xpath('./p:spPr/a:effectLst/*')),0)

    def test_import_keeps_default_single_line_spacing(self):
        source=Presentation()
        slide=source.slides.add_slide(source.slide_layouts[6])
        title=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2))
        title.text='First line\nSecond line'
        data=io.BytesIO();source.save(data)
        project=convert_pptx(data.getvalue(),'single-spacing.pptx')
        self.assertEqual(project['slides'][0]['objects'][0]['style']['lineHeight'],1.0)
        exported=Presentation(io.BytesIO(project_to_pptx(project)))
        self.assertEqual(exported.slides[0].shapes[0].text_frame.paragraphs[0].line_spacing,1.0)

    def test_native_shapes_unicode_text_and_spacing(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        objects=[
            {'id':'arch','type':'shape','name':'Arch','kind':'oval','x':45,'y':70,'w':160,'h':200,'z':1,'style':{'fill':'#1D5D46','borderColor':'#D5DBC1','borderWidth':1.5,'radius':999}},
            {'id':'outline','type':'shape','name':'Outline','kind':'rectangle','x':300,'y':340,'w':190,'h':45,'z':2,'style':{'fill':'transparent','borderColor':'#1D5D46','borderWidth':2}},
            {'id':'title','type':'text','name':'Heading','content':'한글도 편집합니다.\nA little space.','x':260,'y':80,'w':630,'h':150,'z':3,'style':{'fontFamily':'Georgia, serif','fontSize':48,'fontWeight':400,'color':'#303B38','lineHeight':1.12,'padding':0}}]
        project={'name':'Editable sample','slides':[{'title':'Brand story','width':960,'height':540,'background':'#F5F3EA','notes':'Editable presentation sample.','objects':objects}]}
        data=project_to_pptx(project);native=Presentation(io.BytesIO(data));self.assertEqual(native.slides[0].shapes[0].auto_shape_type,MSO_AUTO_SHAPE_TYPE.OVAL)
        reopened=convert_pptx(data,'sample.pptx');self.assertEqual(reopened['name'],project['name']);slide=reopened['slides'][0]
        self.assertEqual(slide['title'],'Brand story');self.assertEqual(slide['background'],'#F5F3EA');self.assertIn('Editable presentation sample.',slide['notes'])
        self.assertEqual(len(slide['objects']),len(objects))
        for before,after in zip(objects,slide['objects']):
            self.assertEqual(before['name'],after['name']);self.assertEqual(before['type'],after['type'])
            for key in ['x','y','w','h']:self.assertAlmostEqual(before[key],after[key],delta=.15)
            if before['type']=='shape':
                self.assertEqual(before['kind'],after['kind']);self.assertEqual(before['style']['fill'],after['style']['fill']);self.assertAlmostEqual(before['style']['borderWidth'],after['style']['borderWidth'],delta=.02)
            else:
                self.assertEqual(before['content'],after['content']);self.assertAlmostEqual(before['style']['fontSize'],after['style']['fontSize'],delta=.1);self.assertEqual(before['style']['lineHeight'],after['style']['lineHeight']);self.assertEqual(after['style']['padding'],0)

    def test_image_remains_an_embedded_editable_picture(self):
        from PIL import Image
        image=Image.new('RGB',(12,12),'#1D5D46');buf=io.BytesIO();image.save(buf,format='PNG')
        source='data:image/png;base64,'+base64.b64encode(buf.getvalue()).decode()
        project={'slides':[{'width':960,'height':540,'objects':[{'type':'image','name':'Embedded artwork','x':50,'y':60,'w':120,'h':120,'src':source,'style':{}}]}]}
        result=convert_pptx(project_to_pptx(project),'picture.pptx')['slides'][0]['objects'][0]
        self.assertEqual(result['type'],'image');self.assertEqual(result['src'],source)

    def test_design_rules_and_invalid_input(self):
        from pathlib import Path
        result=parse_design_md(Path(__file__).with_name('DESIGN.md').read_bytes(),'DESIGN.md')
        self.assertEqual(result['spec']['frontmatter']['colors']['primary'],'#1D5D46');self.assertEqual(len(result['rules']),3)
        for bad in [b'---\nx: &x [1]\ny: *x\n---\n',b'\x00',b'']:
            with self.assertRaises(ValueError):parse_design_md(bad,'DESIGN.md')

    def test_rejects_empty_presentations(self):
        with self.assertRaises(ValueError):dispatch_json(json.dumps({'operation':'export','project':{'slides':[]}}))

if __name__=='__main__':unittest.main()
